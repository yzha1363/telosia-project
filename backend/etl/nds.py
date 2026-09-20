"""
Cleans the Safe Work Australia National Dataset for Compensation Based Statistics.

This ETL currently keeps the two dimensions used by Telosia:
- industry (Tables 1.1 to 1.5)
- occupation (Tables 2.1 to 2.5)

Occupation claims are published at ANZSCO Major Group and Sub-major Group level.
Industry claims are published at ANZSIC Division and Subdivision level.

Output:
    data/processed/nds_category_clean.csv
    data/processed/nds_claims_clean.csv

Run:
    python -m etl.nds
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAW_FILE = (
    ROOT
    / "datasets"
    / "National Dataset for Compensation Based Statistics detailed data file.xlsx"
)
CATEGORY_OUTPUT = ROOT / "data" / "processed" / "nds_category_clean.csv"
CLAIMS_OUTPUT = ROOT / "data" / "processed" / "nds_claims_clean.csv"

HEADER_ROW = 7
FIRST_DATA_ROW = 8
FIRST_YEAR_COL = 2
LAST_YEAR_COL = 11

DIMENSIONS = {
    "1": ("industry", "ANZSIC"),
    "2": ("occupation", "ANZSCO"),
}

MEASURES = {
    "1": ("claim_count", "claims"),
    "2": ("frequency_rate", "claims per million hours worked"),
    "3": ("incidence_rate", "claims per 1,000 employees"),
    "4": ("median_compensation", "AUD"),
    "5": ("median_time_lost", "weeks"),
}

NOT_PUBLISHED = {"NP", "N.P.", "N/P"}
JUNK_CODES = {"&", "&&", "@", "@@"}

EXPECTED_CATEGORY_ROWS = 186
EXPECTED_INDUSTRY_CATEGORIES = 126
EXPECTED_OCCUPATION_CATEGORIES = 60
EXPECTED_CLAIM_ROWS = 9300
EXPECTED_NOT_PUBLISHED = 1705
EXPECTED_GENUINE_BLANKS = 2
EXPECTED_PRELIMINARY_ROWS = 930
EXPECTED_YEARS = 10

CATEGORY_COLUMNS = [
    "dimension",
    "classification",
    "category_code",
    "category_label",
    "level",
    "parent_code",
    "is_nfd",
]

CLAIM_COLUMNS = [
    "dimension",
    "category_code",
    "financial_year",
    "measure",
    "unit",
    "value",
    "is_not_published",
    "is_preliminary",
    "source_sheet",
]

log = logging.getLogger(__name__)


class NdsValidationError(Exception):
    """Raised when the cleaned NDS data fails validation."""


def split_code_label(value: object) -> tuple[str | None, str]:
    """Split a source category such as '42 Carers and aides'."""
    text = str(value).strip()

    if text.lower() == "total":
        return "TOTAL", "Total"

    match = re.match(r"^([A-Za-z0-9&@]{1,3})\s+(.*)$", text)

    if not match:
        return None, text

    return match.group(1), match.group(2).strip()


def normalise_year(value: object) -> tuple[str, bool]:
    """Return a standard financial year and whether it is preliminary."""
    text = str(value).strip()
    is_preliminary = text.lower().endswith("p")

    if is_preliminary:
        text = text[:-1]

    if not re.fullmatch(r"\d{4}-\d{2}", text):
        raise NdsValidationError(
            f"Unexpected financial year value: {value}"
        )

    return text, is_preliminary


def clean_value(value: object) -> tuple[float | None, bool]:
    """Keep NP and genuine blanks separate."""
    if value is None or pd.isna(value):
        return None, False

    if isinstance(value, str):
        text = value.strip()

        if text.upper() in NOT_PUBLISHED:
            return None, True

        if not text:
            return None, False

        try:
            return float(
                text
                .replace(",", "")
                .replace("$", "")
            ), False
        except ValueError as exc:
            raise NdsValidationError(
                f"Unexpected NDS value: {value}"
            ) from exc

    return float(value), False


def classify_category(
    dimension: str,
    code: str,
    label: str,
    current_parent: str | None,
) -> tuple[str, str | None, bool, str | None]:
    """Classify a category and return its level, parent and NFD flag."""
    if code == "TOTAL":
        return "total", None, False, current_parent

    lower_label = label.lower()
    is_nfd = "nfd" in lower_label

    if dimension == "industry":
        if len(code) == 1 and code.isalpha():
            return "division", None, False, code

        if (
            len(code) == 2
            and code[0].isalpha()
            and code[1] == "0"
        ):
            return "division_nfd", code[0], True, current_parent

        if code.isdigit() and len(code) == 2:
            parent = current_parent

            if code == "99":
                parent = None

            return "subdivision", parent, is_nfd, current_parent

    if dimension == "occupation":
        if code.isdigit() and len(code) == 1:
            return "major_group", None, False, code

        if code.isdigit() and len(code) == 2:
            level = (
                "sub_major_group_nfd"
                if code.endswith("0") or is_nfd
                else "sub_major_group"
            )

            return level, code[0], is_nfd, current_parent

    raise NdsValidationError(
        f"Could not classify {dimension} category {code} {label}"
    )


def read_nds_sheet(
    worksheet,
    dimension: str,
    classification: str,
    measure: str,
    unit: str,
) -> tuple[list[dict], list[dict]]:
    """Convert one NDS table to category and claim rows."""
    years = []

    for column in range(
        FIRST_YEAR_COL,
        LAST_YEAR_COL + 1,
    ):
        year, preliminary = normalise_year(
            worksheet.cell(
                HEADER_ROW,
                column,
            ).value
        )

        years.append(
            (
                year,
                preliminary,
            )
        )

    categories = []
    claims = []
    current_parent = None

    for row_number in range(
        FIRST_DATA_ROW,
        worksheet.max_row + 1,
    ):
        raw_category = worksheet.cell(
            row_number,
            1,
        ).value

        if raw_category is None:
            continue

        raw_category = str(
            raw_category
        ).strip()

        if not raw_category:
            continue

        code, label = split_code_label(
            raw_category
        )

        if code in JUNK_CODES:
            continue

        if code is None:
            raise NdsValidationError(
                f"Could not read category on {worksheet.title}: "
                f"{raw_category}"
            )

        (
            level,
            parent_code,
            is_nfd,
            current_parent,
        ) = classify_category(
            dimension,
            code,
            label,
            current_parent,
        )

        categories.append(
            {
                "dimension": dimension,
                "classification": classification,
                "category_code": code,
                "category_label": label,
                "level": level,
                "parent_code": parent_code,
                "is_nfd": is_nfd,
            }
        )

        for column, (
            year,
            preliminary,
        ) in zip(
            range(
                FIRST_YEAR_COL,
                LAST_YEAR_COL + 1,
            ),
            years,
        ):
            value, is_not_published = clean_value(
                worksheet.cell(
                    row_number,
                    column,
                ).value
            )

            claims.append(
                {
                    "dimension": dimension,
                    "category_code": code,
                    "financial_year": year,
                    "measure": measure,
                    "unit": unit,
                    "value": value,
                    "is_not_published": is_not_published,
                    "is_preliminary": preliminary,
                    "source_sheet": worksheet.title,
                }
            )

    return categories, claims


def extract_nds(
    path: Path = RAW_FILE,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Read the industry and occupation tables from the NDS workbook."""
    if not path.exists():
        raise FileNotFoundError(
            f"NDS workbook not found at {path}. "
            "Place the downloaded file in datasets/ without renaming it."
        )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    all_categories = []
    all_claims = []

    for worksheet in workbook.worksheets:
        if not worksheet.title.startswith(
            "Table "
        ):
            continue

        match = re.fullmatch(
            r"Table (\d+)\.(\d+)",
            worksheet.title,
        )

        if not match:
            continue

        (
            dimension_number,
            measure_number,
        ) = match.groups()

        if dimension_number not in DIMENSIONS:
            continue

        if measure_number not in MEASURES:
            raise NdsValidationError(
                f"Unexpected NDS measure sheet: {worksheet.title}"
            )

        dimension, classification = DIMENSIONS[
            dimension_number
        ]

        measure, unit = MEASURES[
            measure_number
        ]

        categories, claims = read_nds_sheet(
            worksheet,
            dimension,
            classification,
            measure,
            unit,
        )

        all_categories.extend(
            categories
        )

        all_claims.extend(
            claims
        )

    categories = pd.DataFrame(
        all_categories,
        columns=CATEGORY_COLUMNS,
    )

    if categories.empty:
        raise NdsValidationError(
            "No NDS industry or occupation categories were extracted."
        )

    conflicts = (
        categories
        .groupby(
            [
                "dimension",
                "category_code",
            ],
            dropna=False,
        )[
            [
                "classification",
                "category_label",
                "level",
                "parent_code",
                "is_nfd",
            ]
        ]
        .nunique(
            dropna=False
        )
    )

    bad_conflicts = conflicts[
        (
            conflicts > 1
        ).any(
            axis=1
        )
    ]

    if not bad_conflicts.empty:
        raise NdsValidationError(
            "A category has inconsistent metadata across NDS sheets: "
            f"{bad_conflicts.index.tolist()[:10]}"
        )

    categories = (
        categories
        .drop_duplicates(
            [
                "dimension",
                "category_code",
            ]
        )
        .sort_values(
            [
                "dimension",
                "category_code",
            ],
            kind="stable",
        )
        .reset_index(
            drop=True
        )
    )

    claims = (
        pd.DataFrame(
            all_claims,
            columns=CLAIM_COLUMNS,
        )
        .sort_values(
            [
                "dimension",
                "category_code",
                "financial_year",
                "measure",
            ],
            kind="stable",
        )
        .reset_index(
            drop=True
        )
    )

    log.info(
        "Extracted %d categories and %d claim rows",
        len(categories),
        len(claims),
    )

    return categories, claims


def validate_nds(
    categories: pd.DataFrame,
    claims: pd.DataFrame,
) -> None:
    """Run structural and value checks on the cleaned NDS data."""
    failures: list[str] = []

    def check(
        label: str,
        condition: bool,
        detail: str = "",
    ) -> None:
        if condition:
            log.info(
                "  PASS  %s",
                label,
            )
        else:
            log.error(
                "  FAIL  %s%s",
                label,
                f" - {detail}" if detail else "",
            )

            failures.append(
                label
            )

    log.info(
        "Validating NDS dataset"
    )

    check(
        f"category rows == {EXPECTED_CATEGORY_ROWS}",
        len(categories) == EXPECTED_CATEGORY_ROWS,
        f"got {len(categories)}",
    )

    check(
        f"claim rows == {EXPECTED_CLAIM_ROWS}",
        len(claims) == EXPECTED_CLAIM_ROWS,
        f"got {len(claims)}",
    )

    industry_categories = categories[
        categories["dimension"] == "industry"
    ]

    occupation_categories = categories[
        categories["dimension"] == "occupation"
    ]

    check(
        f"industry categories == {EXPECTED_INDUSTRY_CATEGORIES}",
        len(industry_categories)
        == EXPECTED_INDUSTRY_CATEGORIES,
        f"got {len(industry_categories)}",
    )

    check(
        f"occupation categories == {EXPECTED_OCCUPATION_CATEGORIES}",
        len(occupation_categories)
        == EXPECTED_OCCUPATION_CATEGORIES,
        f"got {len(occupation_categories)}",
    )

    check(
        "(dimension, category_code) is unique",
        not categories.duplicated(
            [
                "dimension",
                "category_code",
            ]
        ).any(),
    )

    check(
        "claim business key is unique",
        not claims.duplicated(
            [
                "dimension",
                "category_code",
                "financial_year",
                "measure",
            ]
        ).any(),
    )

    check(
        "only industry and occupation were extracted",
        set(
            categories["dimension"]
        )
        == {
            "industry",
            "occupation",
        },
    )

    check(
        f"financial years == {EXPECTED_YEARS}",
        claims["financial_year"].nunique()
        == EXPECTED_YEARS,
        f"got {claims['financial_year'].nunique()}",
    )

    check(
        "all five measures are present",
        set(
            claims["measure"]
        )
        == {
            measure
            for measure, _ in MEASURES.values()
        },
    )

    not_published = int(
        claims[
            "is_not_published"
        ].sum()
    )

    check(
        f"not-published rows == {EXPECTED_NOT_PUBLISHED}",
        not_published
        == EXPECTED_NOT_PUBLISHED,
        f"got {not_published}",
    )

    genuine_blanks = int(
        (
            claims["value"].isna()
            & ~claims["is_not_published"]
        ).sum()
    )

    check(
        f"genuine blank rows == {EXPECTED_GENUINE_BLANKS}",
        genuine_blanks
        == EXPECTED_GENUINE_BLANKS,
        f"got {genuine_blanks}",
    )

    contradictory = claims[
        claims["is_not_published"]
        & claims["value"].notna()
    ]

    check(
        "NP rows do not contain values",
        contradictory.empty,
        f"{len(contradictory)} contradictory rows",
    )

    preliminary = claims[
        claims["is_preliminary"]
    ]

    check(
        f"preliminary rows == {EXPECTED_PRELIMINARY_ROWS}",
        len(preliminary)
        == EXPECTED_PRELIMINARY_ROWS,
        f"got {len(preliminary)}",
    )

    check(
        "only 2023-24 is preliminary",
        set(
            preliminary["financial_year"]
        )
        == {
            "2023-24"
        },
    )

    numeric_values = claims[
        "value"
    ].dropna()

    check(
        "published values are non-negative",
        bool(
            (
                numeric_values >= 0
            ).all()
        ),
        (
            f"min {numeric_values.min()}"
            if not numeric_values.empty
            else "no published values"
        ),
    )

    occupation_submajor = categories[
        (
            categories["dimension"]
            == "occupation"
        )
        & (
            categories["level"]
            == "sub_major_group"
        )
    ]

    check(
        "occupation sub-major groups use 2-digit codes",
        bool(
            occupation_submajor[
                "category_code"
            ]
            .str.fullmatch(
                r"\d{2}"
            )
            .all()
        ),
    )

    carers = categories[
        (
            categories["dimension"]
            == "occupation"
        )
        & (
            categories["category_code"]
            == "42"
        )
    ]

    check(
        "ANZSCO 42 Carers and aides is present",
        len(carers) == 1
        and carers.iloc[0]["category_label"]
        == "Carers and aides"
        and carers.iloc[0]["level"]
        == "sub_major_group"
        and carers.iloc[0]["parent_code"]
        == "4",
    )

    for dimension, expected_categories in (
        (
            "industry",
            EXPECTED_INDUSTRY_CATEGORIES,
        ),
        (
            "occupation",
            EXPECTED_OCCUPATION_CATEGORIES,
        ),
    ):
        for measure, _ in MEASURES.values():
            count = len(
                claims[
                    (
                        claims["dimension"]
                        == dimension
                    )
                    & (
                        claims["measure"]
                        == measure
                    )
                ]
            )

            expected = (
                expected_categories
                * EXPECTED_YEARS
            )

            check(
                f"{dimension} {measure} rows == {expected}",
                count == expected,
                f"got {count}",
            )

    if failures:
        raise NdsValidationError(
            f"{len(failures)} validation check(s) failed:\n  - "
            + "\n  - ".join(
                failures
            )
        )

    log.info(
        "All validation checks passed"
    )


def main(
    write_csv: bool = True,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    logging.basicConfig(
        level=logging.INFO,
        format="%(levelname)-7s %(message)s",
    )

    log.info(
        "Reading %s",
        RAW_FILE.name,
    )

    categories, claims = extract_nds(
        RAW_FILE
    )

    validate_nds(
        categories,
        claims,
    )

    if write_csv:
        CATEGORY_OUTPUT.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        categories.to_csv(
            CATEGORY_OUTPUT,
            index=False,
        )

        claims.to_csv(
            CLAIMS_OUTPUT,
            index=False,
        )

        log.info(
            "Wrote %s (%d rows)",
            CATEGORY_OUTPUT,
            len(categories),
        )

        log.info(
            "Wrote %s (%d rows)",
            CLAIMS_OUTPUT,
            len(claims),
        )

    return categories, claims


if __name__ == "__main__":
    main()
