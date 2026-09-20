from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

SOURCE_FILE = (
    ROOT
    / "data"
    / "annotation"
    / "occupational_hazard_body_part_annotations_all_57_english.xlsx"
)

OUTPUT_FILE = ROOT / "knowledge" / "rag" / "documents.jsonl"
BODY_HAZARDS_PER_CHUNK = 8
HAZARD_BODY_MAPPINGS_PER_CHUNK = 4


def read_sheet(workbook, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]

    records = []

    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue

        records.append(dict(zip(headers, values)))

    return records


def normalise_identifier(value: Any) -> str:
    """Prevent Excel number 1 and text '1' from becoming different IDs."""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def clean(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def build_methodology_documents(
    methodology: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Split methodology into short, topic-specific reference chunks."""

    grouped_lines: dict[str, list[str]] = defaultdict(list)
    grouped_urls: dict[str, list[str]] = defaultdict(list)

    for row in methodology:
        field = clean(row["field"])
        content = clean(row["content"])

        if field.startswith("Score ") or field == "Main-table score meaning":
            group = "scoring"
        elif field in {
            "Normalized mapping purpose",
            "Pair-score continuity",
            "N/A handling",
        }:
            group = "mapping"
        elif (
            field == "Data basis"
            or field.startswith("Important limitation")
        ):
            group = "limitations"
        elif (
            field == "Quality status"
            or field.startswith("Primary source")
        ):
            group = "quality-and-sources"
        else:
            group = "scope"

        if content.startswith(("http://", "https://")):
            grouped_urls[group].append(content)
        else:
            grouped_lines[group].append(f"{field}: {content}")

    documents = []
    for group, lines in grouped_lines.items():
        documents.append(
            {
                "chunk_id": f"methodology-{group}",
                "document_type": "methodology",
                "text": "\n".join(lines),
                "metadata": {
                    "methodology_section": group,
                    "source_urls": grouped_urls.get(group, []),
                    "source_file": SOURCE_FILE.name,
                    "review_status": "reference",
                },
            }
        )

    return documents


def build_documents(include_unreviewed: bool) -> list[dict[str, Any]]:
    workbook = load_workbook(
        SOURCE_FILE,
        data_only=True,
        read_only=True,
    )

    annotations = read_sheet(workbook, "Hazard_Annotations")
    body_mappings = read_sheet(workbook, "Hazard_Body_Map")
    toocs_rows = read_sheet(workbook, "TOOCS_Codebook")
    methodology = read_sheet(workbook, "Methodology")

    toocs_by_code = {
        normalise_identifier(row["TOOCS_3_2_code"]): {
            "body_part": clean(row["title_en"]),
            "source_url": clean(row["source_url"]),
            "use_note": clean(row["use_note_en"]),
        }
        for row in toocs_rows
        if normalise_identifier(row["TOOCS_3_2_code"]) != "N/A"
    }

    mappings_by_hazard: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for mapping in body_mappings:
        hazard_id = normalise_identifier(mapping["hazard_variable_id"])
        body_code = normalise_identifier(
            mapping["body_part_code_TOOCS_3_2"]
        )

        score = int(mapping["association_score_0_to_3"])
        if score not in {0, 1, 2, 3}:
            raise ValueError(
                f"Invalid association score {score} for hazard {hazard_id}."
            )

        if body_code != "N/A" and body_code not in toocs_by_code:
            raise ValueError(
                f"Unknown TOOCS body code {body_code} for hazard {hazard_id}."
            )

        if body_code == "N/A" and score != 0:
            raise ValueError(
                f"N/A mapping for hazard {hazard_id} must have score 0."
            )

        mappings_by_hazard[hazard_id].append(
            {
                "body_code": body_code,
                "body_part": clean(mapping["body_part_en"]),
                "association_score": score,
                "strength": clean(mapping["strength_label_en"]),
                "rationale": clean(mapping["rationale_en"]),
            }
        )

    allowed_statuses = {"approved"}

    if include_unreviewed:
        allowed_statuses.add("needs_human_review")

    documents = []
    included_annotations: dict[str, dict[str, Any]] = {}

    for annotation in annotations:
        hazard_id = normalise_identifier(
            annotation["hazard_variable_id"]
        )
        review_status = clean(annotation["review_status"]).lower()

        if review_status not in allowed_statuses:
            continue

        included_annotations[hazard_id] = annotation

        body_mappings_for_hazard = mappings_by_hazard.get(
            hazard_id,
            [],
        )

        if not body_mappings_for_hazard:
            raise ValueError(
                f"No body mapping found for hazard {hazard_id}."
            )

        direct_mappings = [
            mapping
            for mapping in body_mappings_for_hazard
            if mapping["body_code"] != "N/A"
        ]
        mapping_groups = (
            [
                direct_mappings[start:start + HAZARD_BODY_MAPPINGS_PER_CHUNK]
                for start in range(
                    0,
                    len(direct_mappings),
                    HAZARD_BODY_MAPPINGS_PER_CHUNK,
                )
            ]
            if direct_mappings
            else [[]]
        )
        part_count = len(mapping_groups)

        for part_number, mapping_group in enumerate(mapping_groups, start=1):
            body_text_lines = [
                (
                    f'- {mapping["body_part"]} '
                    f'(TOOCS {mapping["body_code"]}): '
                    f'{mapping["strength"]}, association score '
                    f'{mapping["association_score"]} out of 3.'
                )
                for mapping in mapping_group
            ]
            if not body_text_lines:
                body_text_lines.append(
                    "No defensible direct body-part mapping is available."
                )
            body_text = "\n".join(body_text_lines)

            text = f"""
Hazard variable: {clean(annotation["hazard_variable"])}
Hazard category: {clean(annotation["hazard_category_en"])}

Description:
{clean(annotation["description_en"])}

Potential body-part associations:
{body_text}

Overall association score:
{clean(annotation["association_score_0_to_3"])} out of 3.

Evidence:
{clean(annotation["evidence_title"])}

Evidence summary:
{clean(annotation["evidence_summary"])}

Annotation notes:
{clean(annotation["annotation_notes"])}

Review status:
{review_status}

Important limitation:
This is a modelled exposure-to-potential-body-part association.
It is not an observed injury claim, clinical diagnosis, or proof
that the hazard caused an injury.
""".strip()

            chunk_id = f"hazard-{hazard_id}"
            if part_count > 1:
                chunk_id = f"{chunk_id}-{part_number}"

            documents.append(
                {
                    "chunk_id": chunk_id,
                    "document_type": "hazard_body_association",
                    "text": text,
                    "metadata": {
                        "hazard_variable_id": hazard_id,
                        "hazard_variable": clean(
                            annotation["hazard_variable"]
                        ),
                        "hazard_category": clean(
                            annotation["hazard_category_en"]
                        ),
                        "part_number": part_number,
                        "part_count": part_count,
                        "body_codes": [
                            mapping["body_code"]
                            for mapping in mapping_group
                        ],
                        "body_mappings": [
                            {
                                "body_code": mapping["body_code"],
                                "body_part": mapping["body_part"],
                                "association_score": mapping[
                                    "association_score"
                                ],
                                "strength": mapping["strength"],
                                "rationale": mapping["rationale"],
                            }
                            for mapping in mapping_group
                        ],
                        "review_status": review_status,
                        "data_type": "modelled_association",
                        "source_urls": [
                            clean(annotation["evidence_url_1"]),
                            clean(annotation["evidence_url_2"]),
                        ],
                        "source_file": SOURCE_FILE.name,
                    },
                }
            )

    body_to_hazards: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for hazard_id, mappings in mappings_by_hazard.items():
        annotation = included_annotations.get(hazard_id)
        if annotation is None:
            continue

        for mapping in mappings:
            body_code = mapping["body_code"]
            if body_code == "N/A":
                continue

            body_to_hazards[body_code].append(
                {
                    "hazard_variable_id": hazard_id,
                    "hazard_variable": clean(
                        annotation["hazard_variable"]
                    ),
                    "hazard_category": clean(
                        annotation["hazard_category_en"]
                    ),
                    "association_score": mapping[
                        "association_score"
                    ],
                    "strength": mapping["strength"],
                    "rationale": mapping["rationale"],
                    "review_status": clean(
                        annotation["review_status"]
                    ).lower(),
                    "source_urls": [
                        clean(annotation["evidence_url_1"]),
                        clean(annotation["evidence_url_2"]),
                    ],
                }
            )

    for body_code, related_hazards in body_to_hazards.items():
        codebook_entry = toocs_by_code[body_code]
        body_part = codebook_entry["body_part"]
        related_hazards.sort(
            key=lambda item: (
                -item["association_score"],
                item["hazard_variable"],
            )
        )

        for start in range(
            0,
            len(related_hazards),
            BODY_HAZARDS_PER_CHUNK,
        ):
            subset = related_hazards[
                start:start + BODY_HAZARDS_PER_CHUNK
            ]
            part_number = start // BODY_HAZARDS_PER_CHUNK + 1

            hazard_lines = "\n".join(
                (
                    f'- {item["hazard_variable"]}: '
                    f'{item["strength"]}, association score '
                    f'{item["association_score"]} out of 3.'
                )
                for item in subset
            )

            source_urls = sorted(
                {
                    url
                    for item in subset
                    for url in item["source_urls"]
                    if url
                }
                | (
                    {codebook_entry["source_url"]}
                    if codebook_entry["source_url"]
                    else set()
                )
            )

            text = f"""
Body part: {body_part}
TOOCS bodily location code: {body_code}

Related occupational hazard and work-demand variables:
{hazard_lines}

Classification note:
{codebook_entry["use_note"]}

Important limitation:
These are modelled potential associations between occupational exposure
variables and body locations. They are not observed injury claim counts,
diagnoses, or proof of causation.
""".strip()

            documents.append(
                {
                    "chunk_id": (
                        f"body-part-{body_code}-{part_number}"
                    ),
                    "document_type": "body_part_hazards",
                    "text": text,
                    "metadata": {
                        "body_code": body_code,
                        "body_part": body_part,
                        "part_number": part_number,
                        "related_hazards": subset,
                        "source_urls": source_urls,
                        "source_file": SOURCE_FILE.name,
                        "data_type": "modelled_association",
                    },
                }
            )

    documents.extend(build_methodology_documents(methodology))

    return documents


def main() -> None:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--include-unreviewed",
        action="store_true",
        help="Include needs_human_review records for local testing.",
    )

    args = parser.parse_args()

    documents = build_documents(
        include_unreviewed=args.include_unreviewed
    )

    hazard_count = sum(
        document["document_type"] == "hazard_body_association"
        for document in documents
    )
    hazard_variable_count = len(
        {
            document["metadata"]["hazard_variable_id"]
            for document in documents
            if document["document_type"] == "hazard_body_association"
        }
    )
    body_chunk_count = sum(
        document["document_type"] == "body_part_hazards"
        for document in documents
    )
    methodology_count = sum(
        document["document_type"] == "methodology"
        for document in documents
    )

    if hazard_count == 0:
        raise RuntimeError(
            "No approved hazard records were found. "
            "For local testing, use --include-unreviewed."
        )

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    with OUTPUT_FILE.open("w", encoding="utf-8") as output:
        for document in documents:
            output.write(
                json.dumps(document, ensure_ascii=False) + "\n"
            )

    print(f"Created: {OUTPUT_FILE}")
    print(f"Hazard variables: {hazard_variable_count}")
    print(f"Hazard chunks: {hazard_count}")
    print(f"Body-part chunks: {body_chunk_count}")
    print(f"Methodology chunks: {methodology_count}")
    print(f"Total chunks: {len(documents)}")


if __name__ == "__main__":
    main()
